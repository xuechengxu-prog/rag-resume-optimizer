# FastAPI后端 —— 企业级RAG服务
import os
import io
import shutil
import tempfile
from typing import Optional
from fastapi import FastAPI, File, UploadFile, Form, HTTPException
from fastapi.responses import JSONResponse, FileResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from contextlib import asynccontextmanager

from rag_pipeline import ResumeOptimizationRAG
from evaluation import RAGEvaluator
from config import config


# 全局RAG实例
rag_system = None
evaluator = None

# 缓存最近一次优化结果（用于下载）
_last_result = {"optimized_resume": "", "report": ""}


@asynccontextmanager
async def lifespan(app: FastAPI):
    global rag_system, evaluator
    rag_system = ResumeOptimizationRAG()
    evaluator = RAGEvaluator()
    yield
    # 清理
    if os.path.exists("./temp"):
        shutil.rmtree("./temp")


app = FastAPI(
    title="企业级简历优化RAG系统",
    description="基于Query改写、混合检索、MMR重排的企业级RAG系统",
    version="2.0.0",
    lifespan=lifespan
)

# CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 挂载静态文件
if os.path.exists("./static"):
    app.mount("/static", StaticFiles(directory="static"), name="static")


class OptimizeRequest(BaseModel):
    optimization_level: str = "standard"
    jd_text: Optional[str] = None


class EvaluateRequest(BaseModel):
    query: str
    response: str
    contexts: list


def save_upload_file(upload_file: UploadFile) -> str:
    """保存上传文件到临时目录"""
    os.makedirs("./temp", exist_ok=True)
    file_path = f"./temp/{upload_file.filename}"

    # 检查文件大小
    content = upload_file.file.read()
    if len(content) > config.MAX_FILE_SIZE:
        raise HTTPException(413, "文件超过50MB限制")

    with open(file_path, "wb") as f:
        f.write(content)

    return file_path


@app.get("/")
async def root():
    """根路径返回前端页面"""
    if os.path.exists("./static/index.html"):
        return FileResponse("./static/index.html")
    return {"message": "企业级简历优化RAG系统 API", "version": "2.0.0"}


@app.post("/api/upload-resume")
async def upload_resume(file: UploadFile = File(...)):
    """上传简历文件"""
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in config.SUPPORTED_EXTENSIONS:
        raise HTTPException(400, f"不支持的文件格式: {ext}")

    try:
        file_path = save_upload_file(file)
        result = rag_system.ingest_resume(file_path)
        return {
            "success": True,
            "message": "简历上传成功",
            "data": {
                "text_length": len(result["text"]),
                "chunks": result["chunks"],
                "metadata": result["metadata"]
            }
        }
    except Exception as e:
        raise HTTPException(500, f"简历解析失败: {str(e)}")


@app.post("/api/upload-jd")
async def upload_jd(file: UploadFile = File(...)):
    """上传JD文件（支持PDF、图片、TXT等）"""
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in config.SUPPORTED_EXTENSIONS:
        raise HTTPException(400, f"不支持的文件格式: {ext}")

    try:
        file_path = save_upload_file(file)
        result = rag_system.ingest_jd(file_path)
        return {
            "success": True,
            "message": "JD上传成功",
            "data": {
                "text_length": len(result["text"]),
                "metadata": result["metadata"]
            }
        }
    except Exception as e:
        raise HTTPException(500, f"JD解析失败: {str(e)}")


@app.post("/api/optimize")
async def optimize_resume(request: OptimizeRequest):
    """执行简历优化"""
    try:
        # 如果提供了JD文本，直接使用
        if request.jd_text:
            rag_system.set_jd_text(request.jd_text)

        # 匹配分析
        match = rag_system.match_resume_jd()

        # 优化简历
        optimized = rag_system.optimize_resume(
            match["match_result"],
            optimization_level=request.optimization_level
        )

        # 生成报告
        report = rag_system.generate_report(match["match_result"], optimized)

        return {
            "success": True,
            "data": {
                "match_result": match["match_result"],
                "retrieved_chunks": match["retrieved_chunks"],
                "optimized_resume": optimized,
                "report": report
            }
        }
    except Exception as e:
        raise HTTPException(500, f"优化失败: {str(e)}")


@app.post("/api/full-pipeline")
async def full_pipeline(
    resume: UploadFile = File(...),
    jd: UploadFile = File(...),
    optimization_level: str = Form("standard")
):
    """完整流水线：上传简历和JD，一键优化"""
    try:
        resume_path = save_upload_file(resume)
        jd_path = save_upload_file(jd)

        result = rag_system.run_full_pipeline(
            resume_path=resume_path,
            jd_path=jd_path,
            optimization_level=optimization_level
        )

        # 缓存结果用于下载
        _last_result["optimized_resume"] = result["optimized_resume"]
        _last_result["report"] = result["report"]

        # 5. 自动评估
        try:
            eval_scores = evaluator.evaluate_pipeline(
                original_resume=result["resume_info"]["text"],
                optimized_resume=result["optimized_resume"],
                jd_text=result["jd_info"]["text"],
                match_result=result["match_result"],
                retrieved_chunks_count=result["retrieved_chunks"]
            )
        except Exception as e:
            eval_scores = {"error": f"评估失败: {str(e)}"}

        return {
            "success": True,
            "data": {
                "resume_info": {
                    "text": result["resume_info"]["text"],
                    "text_length": len(result["resume_info"]["text"]),
                    "chunks": result["resume_info"]["chunks"]
                },
                "jd_info": {
                    "text": result["jd_info"]["text"],
                    "text_length": len(result["jd_info"]["text"])
                },
                "match_result": result["match_result"],
                "retrieved_chunks": result["retrieved_chunks"],
                "optimized_resume": result["optimized_resume"],
                "report": result["report"],
                "evaluation": eval_scores
            }
        }
    except Exception as e:
        raise HTTPException(500, f"流水线执行失败: {str(e)}")


@app.post("/api/evaluate")
async def evaluate_rag(request: EvaluateRequest):
    """评估RAG结果"""
    try:
        result = evaluator.quick_evaluate(
            query=request.query,
            response=request.response,
            contexts=request.contexts
        )
        return {"success": True, "data": result}
    except Exception as e:
        raise HTTPException(500, f"评估失败: {str(e)}")


@app.get("/api/config")
async def get_config():
    """获取系统配置"""
    return {
        "success": True,
        "data": {
            "model": config.QWEN_MODEL,
            "chunk_size": config.CHUNK_SIZE,
            "top_k": config.RETRIEVER_TOP_K,
            "mmr_lambda": config.MMR_LAMBDA,
            "supported_formats": config.SUPPORTED_EXTENSIONS,
        }
    }


@app.post("/api/download-resume")
async def download_resume(format: str = Form("pdf")):
    """下载优化后的简历，支持pdf和docx格式"""
    import fitz  # PyMuPDF
    from docx import Document
    from docx.shared import Pt, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    text = _last_result.get("optimized_resume", "")
    if not text:
        raise HTTPException(400, "暂无优化结果，请先执行优化")

    lines = [line.strip() for line in text.split("\n") if line.strip()]

    if format == "docx":
        doc = Document()
        # 设置默认字体
        style = doc.styles['Normal']
        font = style.font
        font.name = '宋体'
        font.size = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.5

        for line in lines:
            p = doc.add_paragraph(line)
            # 检测是否是标题行（短于30字且不以句号结尾）
            if len(line) < 30 and not line.endswith(("。", ".", "！", "？", "；", "：")):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(16)
                    run.font.bold = True

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=optimized_resume.docx"}
        )

    else:  # pdf
        pdf_doc = fitz.open()
        # 设置页面大小 A4
        page = pdf_doc.new_page(width=595, height=842)

        # 写入文本
        y_pos = 50
        font_size = 12
        for line in lines:
            # 检测标题行
            if len(line) < 30 and not line.endswith(("。", ".", "！", "？", "；", "：")):
                page.insert_text((72, y_pos), line, fontsize=16, fontname="china-s")
                y_pos += 30
            else:
                # 处理长行自动换行
                remaining = line
                while remaining:
                    fit_chars = int((595 - 144) / (font_size * 0.5))  # 估算每行能放多少字
                    if len(remaining) <= fit_chars:
                        page.insert_text((72, y_pos), remaining, fontsize=font_size, fontname="china-s")
                        y_pos += font_size * 2
                        remaining = ""
                    else:
                        page.insert_text((72, y_pos), remaining[:fit_chars], fontsize=font_size, fontname="china-s")
                        y_pos += font_size * 2
                        remaining = remaining[fit_chars:]

                    if y_pos > 780:
                        page = pdf_doc.new_page(width=595, height=842)
                        y_pos = 50

        buffer = io.BytesIO()
        pdf_doc.save(buffer)
        pdf_doc.close()
        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=optimized_resume.pdf"}
        )


@app.post("/api/download-report")
async def download_report(format: str = Form("pdf")):
    """下载改进建议报告，支持pdf和docx格式"""
    import fitz  # PyMuPDF
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    text = _last_result.get("report", "")
    if not text:
        raise HTTPException(400, "暂无报告结果，请先执行优化")

    lines = [line.strip() for line in text.split("\n") if line.strip()]

    if format == "docx":
        doc = Document()
        style = doc.styles['Normal']
        font = style.font
        font.name = '宋体'
        font.size = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.5

        for line in lines:
            p = doc.add_paragraph(line)
            if len(line) < 30 and not line.endswith(("。", ".", "！", "？", "；", "：")):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(16)
                    run.font.bold = True

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=optimization_report.docx"}
        )

    else:  # pdf
        pdf_doc = fitz.open()
        page = pdf_doc.new_page(width=595, height=842)

        y_pos = 50
        font_size = 12
        for line in lines:
            if len(line) < 30 and not line.endswith(("。", ".", "！", "？", "；", "：")):
                page.insert_text((72, y_pos), line, fontsize=16, fontname="china-s")
                y_pos += 30
            else:
                remaining = line
                while remaining:
                    fit_chars = int((595 - 144) / (font_size * 0.5))
                    if len(remaining) <= fit_chars:
                        page.insert_text((72, y_pos), remaining, fontsize=font_size, fontname="china-s")
                        y_pos += font_size * 2
                        remaining = ""
                    else:
                        page.insert_text((72, y_pos), remaining[:fit_chars], fontsize=font_size, fontname="china-s")
                        y_pos += font_size * 2
                        remaining = remaining[fit_chars:]

                    if y_pos > 780:
                        page = pdf_doc.new_page(width=595, height=842)
                        y_pos = 50

        buffer = io.BytesIO()
        pdf_doc.save(buffer)
        pdf_doc.close()
        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=optimization_report.pdf"}
        )


@app.get("/{path:path}")
async def serve_frontend(path: str = "index.html"):
    """服务前端页面"""
    file_path = f"./static/{path}"
    if os.path.exists(file_path):
        return FileResponse(file_path)
    if os.path.exists("./static/index.html"):
        return FileResponse("./static/index.html")
    return JSONResponse({"message": "前端文件不存在，请访问 /docs 查看API文档"})


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
